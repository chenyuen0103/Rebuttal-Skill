from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "REVIEWS_RAW.md"
OUTPUT = ROOT / "CausalMix_Reviews_Google_Docs_Source.docx"


def set_font(run, size=11, bold=False, italic=False, color="000000"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    specs = {
        "Heading 1": (20, 20, 6, "000000"),
        "Heading 2": (16, 18, 6, "000000"),
        "Heading 3": (14, 16, 4, "434343"),
    }
    for name, (size, before, after, color) in specs.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_labelled_paragraph(doc, label, remainder):
    p = doc.add_paragraph()
    r = p.add_run(label)
    set_font(r, bold=True)
    if remainder:
        r = p.add_run(remainder)
        set_font(r)
    return p


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    configure_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.keep_with_next = True
    r = title.add_run("Reviews for “CausalMix: Defining and Benchmarking Hybrid Causal Discovery”")
    set_font(r, size=26, bold=False)

    # Skip the Markdown packaging lines. Review text remains verbatim.
    start = next(i for i, line in enumerate(lines) if line.startswith("Meta Review of Submission"))
    for line in lines[start:]:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue

        if stripped.startswith("Meta Review of Submission"):
            add_heading(doc, stripped, 1)
        elif stripped in {"Review of paper 3307", "Interesting ablation study with a flawed SFT experiment design", "review"}:
            add_heading(doc, stripped, 1)
        elif stripped in {
            "Metareview:", "Summary:", "Strengths:", "Weaknesses:", "Limitations:",
            "Questions:", "Reproducibility Comments:", "Dataset Comments:",
            "Paper Formatting Concerns:"
        }:
            add_heading(doc, stripped[:-1], 2)
        elif stripped == "Add:":
            continue
        elif any(stripped.startswith(prefix) for prefix in (
            "Rating:", "Confidence:", "Ethical Concerns:", "Reproducibility:",
            "Dataset Assessment:", "Contribution Type Check:", "LLM Policy:",
            "Code Of Conduct Acknowledgement:", "Responsible Reviewing Acknowledgement:"
        )):
            label, remainder = stripped.split(":", 1)
            add_labelled_paragraph(doc, label + ":", remainder)
        else:
            p = doc.add_paragraph()
            r = p.add_run(line)
            set_font(r)

    # Prevent empty trailing paragraphs from producing a blank page.
    while len(doc.paragraphs) > 1 and not doc.paragraphs[-1].text:
        element = doc.paragraphs[-1]._element
        element.getparent().remove(element)

    props = doc.core_properties
    props.title = "Reviews for CausalMix: Defining and Benchmarking Hybrid Causal Discovery"
    props.subject = "NeurIPS 2026 reviews"
    props.author = ""
    props.last_modified_by = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
