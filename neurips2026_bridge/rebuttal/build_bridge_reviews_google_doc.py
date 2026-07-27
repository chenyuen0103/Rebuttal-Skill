from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT.parent / "review.txt"
OUTPUT = ROOT / "BRIDGE_Official_Reviews_Google_Docs_Source.docx"


def set_font(run, size=11, bold=False, italic=False, color="000000"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    specs = {
        "Heading 1": (20, 20, 6, "000000"),
        "Heading 2": (16, 18, 6, "000000"),
        "Heading 3": (14, 16, 4, "434343"),
    }
    for name, (size, before, after, color) in specs.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)
    return paragraph


def add_labelled_paragraph(doc, label, remainder):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(label)
    set_font(run, bold=True)
    if remainder:
        run = paragraph.add_run(remainder)
        set_font(run)
    return paragraph


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("Meta Review of Submission25120"))
    end = next(i for i, line in enumerate(lines[start:], start) if line.startswith("LLM Feedback by Program Chairs"))
    review_lines = lines[start:end]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("Official Reviews for “End-to-End Sequence-Graph Learning for User Behavior Modeling”")
    set_font(run, size=26)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("NeurIPS 2026 · Submission 25120 · Meta-review and three official reviews")
    set_font(run, size=11, color="555555")

    note = doc.add_paragraph()
    run = note.add_run("Scope: ")
    set_font(run, bold=True)
    run = note.add_run("This document preserves the official human reviews. Automated pre-submission PAT feedback is excluded because it is not part of the official review record and evaluated an earlier manuscript version.")
    set_font(run)

    reviewer_count = 0
    skip_metadata_line = False
    for line in review_lines:
        stripped = line.strip()
        if not stripped or stripped == "Add:":
            continue

        if stripped.startswith("Meta Review of Submission25120"):
            add_heading(doc, "Meta-review · Area Chair xm4H", 1)
            skip_metadata_line = True
            continue
        if stripped.startswith("Official Review of Submission25120 by Reviewer"):
            reviewer_count += 1
            reviewer_id = stripped.rsplit("Reviewer ", 1)[-1]
            add_heading(doc, f"Official review · Reviewer {reviewer_id}", 1)
            skip_metadata_line = True
            continue
        if skip_metadata_line and (stripped.startswith("Meta Reviewby") or stripped.startswith("Official Reviewby")):
            skip_metadata_line = False
            continue

        if stripped in {
            "Metareview:", "Summary:", "Strengths And Weaknesses:", "Strengths:",
            "Weaknesses:", "Questions:", "Limitations:", "Paper Formatting Concerns:"
        }:
            add_heading(doc, stripped[:-1], 2)
        elif stripped in {"Quality", "Clarity", "Significance", "Originality"}:
            add_heading(doc, stripped, 3)
        elif any(stripped.startswith(prefix) for prefix in (
            "Contribution Type:", "Quality:", "Clarity:", "Significance:", "Originality:",
            "Rating:", "Confidence:", "Ethical Concerns:", "Code Of Conduct Acknowledgement:",
            "Responsible Reviewing Acknowledgement:"
        )):
            label, remainder = stripped.split(":", 1)
            add_labelled_paragraph(doc, label + ":", remainder)
        else:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(stripped)
            set_font(run)

    while len(doc.paragraphs) > 1 and not doc.paragraphs[-1].text:
        element = doc.paragraphs[-1]._element
        element.getparent().remove(element)

    properties = doc.core_properties
    properties.title = "Official Reviews for End-to-End Sequence-Graph Learning for User Behavior Modeling"
    properties.subject = "NeurIPS 2026 submission 25120 reviews"
    properties.author = ""
    properties.last_modified_by = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
